"""
Excel Parser for Universal Data Layer.
Parses .xlsx (via openpyxl), .xls (via xlrd), and HTML-table-based .xls files with multi-sheet support.
"""
import re
from pathlib import Path
from typing import Any, Dict, List, Tuple
import openpyxl

def parse_excel_file(filepath: Path, max_rows_per_sheet: int = 20000) -> Tuple[Dict[str, List[Dict[str, Any]]], List[str], str]:
    """
    Parses an Excel spreadsheet. Returns (sheets_dict, sheet_names, status).
    """
    if filepath.stat().st_size == 0:
        return {}, [], "EMPTY_FILE_0_BYTES"

    ext = filepath.suffix.lower()
    sheets_data: Dict[str, List[Dict[str, Any]]] = {}
    sheet_names: List[str] = []

    # 1. Try openpyxl for .xlsx
    if ext == ".xlsx":
        try:
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            for sname in sheet_names:
                ws = wb[sname]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    sheets_data[sname] = []
                    continue
                
                # Find header row (first non-empty row)
                header_idx = 0
                while header_idx < len(rows) and all(c is None for c in rows[header_idx]):
                    header_idx += 1

                if header_idx >= len(rows):
                    sheets_data[sname] = []
                    continue

                headers = [str(c).strip() if c is not None else f"col_{j}" for j, c in enumerate(rows[header_idx])]
                sheet_records = []
                for row_vals in rows[header_idx + 1: header_idx + 1 + max_rows_per_sheet]:
                    if all(v is None for v in row_vals):
                        continue
                    rec = {}
                    for k, v in zip(headers, row_vals):
                        if k:
                            rec[k] = v
                    sheet_records.append(rec)
                sheets_data[sname] = sheet_records

            wb.close()
            return sheets_data, sheet_names, "SUCCESS"
        except Exception as e:
            pass

    # 2. Try xlrd for binary .xls
    try:
        import xlrd
        wb = xlrd.open_workbook(filepath, on_demand=True)
        sheet_names = wb.sheet_names()
        for sname in sheet_names:
            ws = wb.sheet_by_name(sname)
            if ws.nrows == 0:
                sheets_data[sname] = []
                continue
            
            headers = [str(ws.cell_value(0, col)).strip() for col in range(ws.ncols)]
            sheet_records = []
            for row_idx in range(1, min(ws.nrows, max_rows_per_sheet + 1)):
                rec = {}
                for col_idx in range(ws.ncols):
                    k = headers[col_idx] or f"col_{col_idx}"
                    rec[k] = ws.cell_value(row_idx, col_idx)
                sheet_records.append(rec)
            sheets_data[sname] = sheet_records

        return sheets_data, sheet_names, "SUCCESS"
    except Exception as e:
        pass

    # 3. Fallback for HTML-table exported as .xls
    try:
        content = ""
        for enc in ["utf-8", "latin-1", "cp1252"]:
            try:
                with open(filepath, "r", encoding=enc, errors="ignore") as f:
                    content = f.read(500000) # Read initial 500KB
                    break
            except Exception:
                continue

        if "<table" in content.lower() or "<tr" in content.lower():
            # Extract simple rows and cells
            rows = re.findall(r"<tr[^>]*>(.*?)</tr>", content, re.DOTALL | re.IGNORECASE)
            if rows:
                parsed_rows = []
                for r in rows[:max_rows_per_sheet]:
                    cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", r, re.DOTALL | re.IGNORECASE)
                    clean_cells = [re.sub(r"<[^>]+>", "", c).strip() for c in cells]
                    if any(clean_cells):
                        parsed_rows.append(clean_cells)
                
                if parsed_rows:
                    headers = parsed_rows[0]
                    records = []
                    for r in parsed_rows[1:]:
                        rec = {headers[i] if i < len(headers) and headers[i] else f"col_{i}": val for i, val in enumerate(r)}
                        records.append(rec)
                    sheets_data["HTML_EXPORT"] = records
                    return sheets_data, ["HTML_EXPORT"], "SUCCESS_HTML_EXPORT"
    except Exception:
        pass

    return {}, [], "PARSING_FAILED"
