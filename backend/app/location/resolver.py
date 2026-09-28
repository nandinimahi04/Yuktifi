from __future__ import annotations

import csv
import difflib
import re
from pathlib import Path
from typing import Any

ALIASES = {
    "state_name": "state", "state": "state",
    "district_name": "district", "district": "district",
    "subdistrict": "subdistrict", "sub_district": "subdistrict",
    "sub_district_name": "subdistrict", "taluka": "subdistrict", "tehsil": "subdistrict",
    "block_name": "block", "development_block": "block", "block": "block",
    "village_name": "village", "village": "village", "town": "village", "town_name": "village",
    "state_code": "state_code", "district_code": "district_code",
    "subdistrict_code": "subdistrict_code", "sub_district_code": "subdistrict_code",
    "block_code": "block_code", "village_code": "village_code", "location_code": "location_code",
    "latitude": "lat", "lat": "lat", "longitude": "lng", "lon": "lng", "lng": "lng",
    "level": "admin_level", "admin_level": "admin_level", "type": "admin_level",
}


def norm(value: Any) -> str:
    value = "" if value is None else str(value)
    value = value.casefold().replace("&", " and ")
    value = re.sub(r"\([^)]*\)", "", value)
    value = re.sub(r"[^\w\s]", " ", value, flags=re.UNICODE)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def _key(header: Any) -> str:
    s = norm(header).replace(" ", "_")
    return ALIASES.get(s, s)


def _canonical_row(row: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for k, v in row.items():
        ck = _key(k)
        if ck:
            out[ck] = v
    for field in ("state", "district", "subdistrict", "block", "village", "admin_level"):
        if out.get(field) is not None:
            out[field] = str(out[field]).strip()
    for field in ("lat", "lng"):
        if out.get(field) not in (None, ""):
            try:
                out[field] = float(out[field])
            except (TypeError, ValueError):
                out.pop(field, None)
    return out


def read_master(path: str | Path) -> list[dict[str, Any]]:
    """Read an LGD/Census-derived master in CSV/XLSX/ZIP form."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    if path.suffix.lower() == ".zip":
        import zipfile
        with zipfile.ZipFile(path) as zf:
            members = [n for n in zf.namelist() if Path(n).suffix.lower() in {".csv", ".xlsx"} and not n.endswith("/")]
            if not members:
                raise ValueError("Location ZIP contains no CSV/XLSX file")
            # Prefer files whose names indicate village/local-government geography.
            members.sort(key=lambda n: (0 if any(x in n.lower() for x in ("village", "district", "subdistrict", "block", "lgd")) else 1, n))
            tmp = Path("/tmp") / f"yukti_location_{Path(members[0]).name}"
            tmp.write_bytes(zf.read(members[0]))
            return read_master(tmp)
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return [_canonical_row(r) for r in csv.DictReader(f)]
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        rows: list[dict[str, Any]] = []
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            try:
                headers = next(it)
            except StopIteration:
                continue
            headers = [_key(h) for h in headers]
            if not any(headers):
                continue
            for vals in it:
                if not any(v is not None for v in vals):
                    continue
                rows.append(_canonical_row(dict(zip(headers, vals))))
        return rows
    raise ValueError("Unsupported location master format; use CSV, XLSX or ZIP")


class LocationResolver:
    """Deterministic location resolver backed by a government-derived master."""

    def __init__(self, rows: list[dict[str, Any]] | None = None):
        self.rows = [r for r in (rows or []) if r.get("state") or r.get("district") or r.get("village")]

    @classmethod
    def from_file(cls, path: str | Path) -> "LocationResolver":
        return cls(read_master(path))

    def _score(self, row: dict[str, Any], query: str) -> float:
        nq = norm(query)
        parts = [row.get("village"), row.get("block"), row.get("subdistrict"), row.get("district"), row.get("state")]
        text = ", ".join(str(x) for x in parts if x)
        nt = norm(text)
        if nq == nt:
            return 1.0
        if nq in nt:
            return 0.93
        qparts = [p for p in re.split(r"[,/|]", query) if norm(p)]
        score = 0.0
        for qp in qparts:
            best = max((difflib.SequenceMatcher(None, norm(qp), norm(x)).ratio() for x in parts if x), default=0.0)
            score += best
        return score / max(len(qparts), 1)

    def resolve_text(self, query: str, threshold: float = 0.78) -> dict[str, Any]:
        if not query or not self.rows:
            return {"status": "INSUFFICIENT_LOCATION_EVIDENCE", "reason": "no_query_or_master"}
        ranked = sorted(((self._score(r, query), r) for r in self.rows), key=lambda x: x[0], reverse=True)
        score, row = ranked[0]
        if score < threshold:
            return {"status": "INSUFFICIENT_LOCATION_EVIDENCE", "query": query, "best_match": row, "match_score": round(score, 3)}
        return self._result(row, query, score, "master_match")

    def resolve_coordinates(self, lat: float, lng: float) -> dict[str, Any]:
        candidates = [r for r in self.rows if r.get("lat") is not None and r.get("lng") is not None]
        if not candidates:
            return {"status": "INSUFFICIENT_LOCATION_EVIDENCE", "reason": "needs_reverse_geocode", "lat": lat, "lng": lng}
        # Equirectangular distance is sufficient for nearest administrative point at this stage.
        import math
        lat_scale = 111.32
        lon_scale = 111.32 * math.cos(math.radians(lat))
        best = min(candidates, key=lambda r: ((float(r["lat"]) - lat) * lat_scale) ** 2 + ((float(r["lng"]) - lng) * lon_scale) ** 2)
        distance_km = math.sqrt(((float(best["lat"]) - lat) * lat_scale) ** 2 + ((float(best["lng"]) - lng) * lon_scale) ** 2)
        return self._result(best, f"{lat},{lng}", 1.0, "nearest_master_coordinate", distance_km=round(distance_km, 3))

    def resolve(self, query: str | None = None, lat: float | None = None, lng: float | None = None) -> dict[str, Any]:
        if lat is not None and lng is not None:
            return self.resolve_coordinates(lat, lng)
        if query:
            match = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*[, ]\s*(-?\d+(?:\.\d+)?)\s*", query)
            if match:
                return self.resolve_coordinates(float(match.group(1)), float(match.group(2)))
            return self.resolve_text(query)
        return {"status": "INSUFFICIENT_LOCATION_EVIDENCE", "reason": "missing_location"}

    @staticmethod
    def _result(row: dict[str, Any], query: str, score: float, method: str, **extra: Any) -> dict[str, Any]:
        return {
            "status": "resolved",
            "query": query,
            "match_score": round(score, 3),
            "method": method,
            "state": row.get("state"),
            "district": row.get("district"),
            "subdistrict": row.get("subdistrict"),
            "block": row.get("block"),
            "village": row.get("village"),
            "codes": {
                k: row.get(k) for k in ("state_code", "district_code", "subdistrict_code", "block_code", "village_code", "location_code") if row.get(k) not in (None, "")
            },
            "lat": row.get("lat"),
            "lng": row.get("lng"),
            "admin_level": row.get("admin_level"),
            **extra,
        }
